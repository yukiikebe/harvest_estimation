from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


WINDOW_ORDER = ("6mo", "9mo", "1year")
FEATURE_ORDER = ("all_indices", "ndvi_only")
MODEL_ORDER = ("cnn", "rnn", "hybrid", "hybrid_infer")

WINDOW_LABELS = {
    "6mo": "6 month",
    "9mo": "9 month",
    "1year": "12 month",
}
FEATURE_LABELS = {
    "all_indices": "NDVI, NDWI, EVI",
    "ndvi_only": "NDVI",
}
MODEL_LABELS = {
    "cnn": "CNN",
    "rnn": "RNN",
    "hybrid": "Hybrid",
    "hybrid_infer": "Hybrid_infer",
}

METRIC_KEYS = ("mae_mean", "mae_start", "mae_end", "mae_window", "iou_mean", "ordered_pct")
LOWER_IS_BETTER = {"mae_mean", "mae_start", "mae_end", "mae_window"}
HIGHER_IS_BETTER = {"iou_mean", "ordered_pct"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Aggregate all-crop CNN/RNN/Hybrid/Hybrid_infer window sweep results into Excel."
    )
    parser.add_argument("--results-root", type=Path, required=True)
    parser.add_argument("--out-xlsx", type=Path, required=True)
    parser.add_argument("--train-year", type=str, default="2022")
    parser.add_argument("--test-year", type=str, default="2023")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows = collect_rows(args.results_root)
    if not rows:
        raise RuntimeError(f"No result rows found under {args.results_root}")

    args.out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(
        args.out_xlsx,
        rows,
        results_root=args.results_root,
        train_year=args.train_year,
        test_year=args.test_year,
    )
    print(f"Wrote all-crops comparison workbook to {args.out_xlsx}")


def collect_rows(results_root: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    rows.extend(collect_train_model_rows(results_root))
    rows.extend(collect_hybrid_infer_rows(results_root))
    return sorted(rows, key=row_sort_key)


def collect_train_model_rows(results_root: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for model in ("cnn", "rnn", "hybrid"):
        model_root = results_root / model
        if not model_root.exists():
            continue
        for metrics_path in sorted(model_root.glob("*/*/*/metrics.json")):
            try:
                window, feature_set, _crop_dir = metrics_path.relative_to(model_root).parts[:3]
            except ValueError:
                continue
            if window not in WINDOW_ORDER or feature_set not in FEATURE_ORDER:
                continue
            with open(metrics_path, "r", encoding="utf-8") as f:
                metrics = json.load(f)
            test_metrics = metrics.get("test_metrics") or {}
            if not test_metrics:
                continue
            row = base_row(
                crop=str(metrics.get("crop", metrics_path.parent.name)),
                model=model,
                window=window,
                feature_set=feature_set,
                source_path=metrics_path,
            )
            row.update(
                {
                    "num_train": metrics.get("num_train"),
                    "num_val": metrics.get("num_val"),
                    "num_test": metrics.get("num_test"),
                    "num_records": metrics.get("num_test"),
                }
            )
            for key in METRIC_KEYS:
                row[key] = to_optional_float(test_metrics.get(key))
            rows.append(row)
    return rows


def collect_hybrid_infer_rows(results_root: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    infer_root = results_root / "hybrid_infer"
    if not infer_root.exists():
        return rows

    for csv_path in sorted(infer_root.glob("*/*/predictions.csv")):
        try:
            window, feature_set, _file_name = csv_path.relative_to(infer_root).parts
        except ValueError:
            continue
        if window not in WINDOW_ORDER or feature_set not in FEATURE_ORDER:
            continue

        grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
        with open(csv_path, "r", newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for csv_row in reader:
                crop = str(csv_row.get("crop", "")).strip()
                if crop:
                    grouped[crop].append(csv_row)

        for crop, crop_rows in grouped.items():
            metrics = compute_prediction_metrics(crop_rows)
            row = base_row(
                crop=crop,
                model="hybrid_infer",
                window=window,
                feature_set=feature_set,
                source_path=csv_path,
            )
            row.update(
                {
                    "num_train": None,
                    "num_val": None,
                    "num_test": metrics["num_records"],
                    "num_records": metrics["num_records"],
                    **{key: metrics.get(key) for key in METRIC_KEYS},
                }
            )
            rows.append(row)
    return rows


def compute_prediction_metrics(rows: list[dict[str, str]]) -> dict[str, float]:
    starts_true: list[float] = []
    ends_true: list[float] = []
    starts_pred: list[float] = []
    ends_pred: list[float] = []
    ious: list[float] = []

    for row in rows:
        true_start = to_optional_float(row.get("true_start_doy"))
        true_end = to_optional_float(row.get("true_end_doy"))
        pred_start = to_optional_float(row.get("pred_start_doy"))
        pred_end = to_optional_float(row.get("pred_end_doy"))
        if None in (true_start, true_end, pred_start, pred_end):
            continue
        starts_true.append(float(true_start))
        ends_true.append(float(true_end))
        starts_pred.append(float(pred_start))
        ends_pred.append(float(pred_end))
        iou = to_optional_float(row.get("window_iou"))
        if iou is not None:
            ious.append(float(iou))

    if not starts_true:
        return {key: None for key in METRIC_KEYS} | {"num_records": 0}

    mae_start = mean(abs(pred - true) for pred, true in zip(starts_pred, starts_true))
    mae_end = mean(abs(pred - true) for pred, true in zip(ends_pred, ends_true))
    mae_window = mean(
        abs((pred_end - pred_start) - (true_end - true_start))
        for pred_start, pred_end, true_start, true_end in zip(
            starts_pred, ends_pred, starts_true, ends_true
        )
    )
    ordered_pct = mean(1.0 if start <= end else 0.0 for start, end in zip(starts_pred, ends_pred))
    return {
        "num_records": len(starts_true),
        "mae_mean": (mae_start + mae_end) / 2.0,
        "mae_start": mae_start,
        "mae_end": mae_end,
        "mae_window": mae_window,
        "iou_mean": mean(ious) if ious else None,
        "ordered_pct": ordered_pct,
    }


def base_row(
    *,
    crop: str,
    model: str,
    window: str,
    feature_set: str,
    source_path: Path,
) -> dict[str, Any]:
    return {
        "crop": crop,
        "model": model,
        "model_label": MODEL_LABELS.get(model, model),
        "window": window,
        "window_label": WINDOW_LABELS.get(window, window),
        "feature_set": feature_set,
        "feature_label": FEATURE_LABELS.get(feature_set, feature_set),
        "source_path": str(source_path),
    }


def write_workbook(
    path: Path,
    rows: list[dict[str, Any]],
    *,
    results_root: Path,
    train_year: str,
    test_year: str,
) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    write_overview_sheet(wb, rows, results_root, train_year, test_year)
    write_best_by_crop_sheet(wb, rows)
    write_model_average_sheet(wb, rows)
    write_all_results_sheet(wb, rows)
    write_crop_tables_sheet(wb, rows)

    wb.save(path)


def write_overview_sheet(
    wb: Workbook,
    rows: list[dict[str, Any]],
    results_root: Path,
    train_year: str,
    test_year: str,
) -> None:
    ws = wb.create_sheet("Overview")
    bold = Font(bold=True)
    overview_rows = [
        ("Results root", str(results_root)),
        ("Train year", train_year),
        ("Test year", test_year),
        ("Result rows", len(rows)),
        ("Crops", len({row["crop"] for row in rows})),
        ("Models", ", ".join(MODEL_LABELS[model] for model in MODEL_ORDER)),
        ("Windows", ", ".join(WINDOW_LABELS[window] for window in WINDOW_ORDER)),
        ("Feature sets", ", ".join(FEATURE_LABELS[feature] for feature in FEATURE_ORDER)),
    ]
    for row_idx, (key, value) in enumerate(overview_rows, start=1):
        ws.cell(row_idx, 1, key).font = bold
        ws.cell(row_idx, 2, value)
    set_widths(ws, {1: 18, 2: 120})


def write_best_by_crop_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Best by Crop")
    headers = [
        "crop",
        "best_mae_model",
        "best_mae_window",
        "best_mae_feature_set",
        "best_mae_mean",
        "best_iou_model",
        "best_iou_window",
        "best_iou_feature_set",
        "best_iou_mean",
        "num_options",
    ]
    write_header(ws, headers)

    rows_by_crop: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        rows_by_crop[row["crop"]].append(row)

    row_idx = 2
    for crop in sorted(rows_by_crop):
        crop_rows = rows_by_crop[crop]
        mae_rows = [row for row in crop_rows if row.get("mae_mean") is not None]
        iou_rows = [row for row in crop_rows if row.get("iou_mean") is not None]
        best_mae = min(mae_rows, key=lambda row: float(row["mae_mean"])) if mae_rows else None
        best_iou = max(iou_rows, key=lambda row: float(row["iou_mean"])) if iou_rows else None
        values = [
            crop,
            label_value(best_mae, "model_label"),
            label_value(best_mae, "window_label"),
            label_value(best_mae, "feature_label"),
            metric_value(best_mae, "mae_mean"),
            label_value(best_iou, "model_label"),
            label_value(best_iou, "window_label"),
            label_value(best_iou, "feature_label"),
            metric_value(best_iou, "iou_mean"),
            len(crop_rows),
        ]
        write_row(ws, row_idx, headers, values)
        row_idx += 1
    format_numeric_columns(ws, first_row=2)
    set_widths(ws, {1: 26, 2: 16, 3: 14, 4: 20, 5: 14, 6: 16, 7: 14, 8: 20, 9: 14, 10: 12})


def write_model_average_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Model Average")
    headers = [
        "model",
        "window",
        "feature_set",
        "crops",
        "avg_records",
        "avg_mae_mean",
        "avg_mae_start",
        "avg_mae_end",
        "avg_mae_window",
        "avg_iou_mean",
        "avg_ordered_pct",
    ]
    write_header(ws, headers)

    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[(row["model"], row["window"], row["feature_set"])].append(row)

    row_idx = 2
    for model in MODEL_ORDER:
        for window in WINDOW_ORDER:
            for feature_set in FEATURE_ORDER:
                group_rows = grouped.get((model, window, feature_set), [])
                if not group_rows:
                    continue
                values = [
                    MODEL_LABELS[model],
                    WINDOW_LABELS[window],
                    FEATURE_LABELS[feature_set],
                    len({row["crop"] for row in group_rows}),
                    average_key(group_rows, "num_records"),
                    average_key(group_rows, "mae_mean"),
                    average_key(group_rows, "mae_start"),
                    average_key(group_rows, "mae_end"),
                    average_key(group_rows, "mae_window"),
                    average_key(group_rows, "iou_mean"),
                    average_key(group_rows, "ordered_pct"),
                ]
                write_row(ws, row_idx, headers, values)
                row_idx += 1
    format_numeric_columns(ws, first_row=2)
    set_widths(ws, {1: 16, 2: 13, 3: 20, 4: 9, 5: 12, 6: 14, 7: 14, 8: 14, 9: 16, 10: 14, 11: 16})


def write_all_results_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("All Results")
    headers = [
        "crop",
        "model",
        "window",
        "feature_set",
        "num_train",
        "num_val",
        "num_test",
        "num_records",
        "mae_mean",
        "mae_start",
        "mae_end",
        "mae_window",
        "iou_mean",
        "ordered_pct",
        "source_path",
    ]
    write_header(ws, headers)
    for row_idx, row in enumerate(rows, start=2):
        values = [
            row["crop"],
            row["model_label"],
            row["window_label"],
            row["feature_label"],
            row.get("num_train"),
            row.get("num_val"),
            row.get("num_test"),
            row.get("num_records"),
            row.get("mae_mean"),
            row.get("mae_start"),
            row.get("mae_end"),
            row.get("mae_window"),
            row.get("iou_mean"),
            row.get("ordered_pct"),
            row.get("source_path"),
        ]
        write_row(ws, row_idx, headers, values)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    format_numeric_columns(ws, first_row=2)
    set_widths(ws, {1: 26, 2: 14, 3: 13, 4: 20, 5: 10, 6: 9, 7: 10, 8: 12, 9: 13, 10: 13, 11: 13, 12: 15, 13: 13, 14: 13, 15: 90})


def write_crop_tables_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Crop Tables")
    bold = Font(bold=True)
    best_fill = PatternFill("solid", fgColor="FFF2CC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    headers = [
        "model",
        "window",
        "feature_set",
        "num_records",
        "mae_mean",
        "mae_start",
        "mae_end",
        "mae_window",
        "iou_mean",
        "ordered_pct",
    ]

    rows_by_crop: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        rows_by_crop[row["crop"]].append(row)

    row_idx = 1
    for crop in sorted(rows_by_crop):
        ws.cell(row_idx, 1, crop).font = bold
        row_idx += 1
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row_idx, col_idx, header)
            cell.font = bold
            cell.fill = header_fill
        row_idx += 1
        first_data_row = row_idx
        for row in sorted(rows_by_crop[crop], key=row_sort_key):
            values = [
                row["model_label"],
                row["window_label"],
                row["feature_label"],
                row.get("num_records"),
                row.get("mae_mean"),
                row.get("mae_start"),
                row.get("mae_end"),
                row.get("mae_window"),
                row.get("iou_mean"),
                row.get("ordered_pct"),
            ]
            write_row(ws, row_idx, headers, values)
            row_idx += 1
        last_data_row = row_idx - 1
        highlight_best_metric_values(ws, first_data_row, last_data_row, headers, bold, best_fill)
        row_idx += 2

    format_numeric_columns(ws, first_row=1)
    set_widths(ws, {1: 16, 2: 13, 3: 20, 4: 12, 5: 13, 6: 13, 7: 13, 8: 15, 9: 13, 10: 13})


def write_header(ws, headers: list[str]) -> None:
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, header)
        cell.font = bold
        cell.fill = fill


def write_row(ws, row_idx: int, headers: list[str], values: list[Any]) -> None:
    for col_idx, value in enumerate(values[: len(headers)], start=1):
        ws.cell(row_idx, col_idx, value)


def highlight_best_metric_values(
    ws,
    first_data_row: int,
    last_data_row: int,
    headers: list[str],
    bold: Font,
    fill: PatternFill,
) -> None:
    if first_data_row > last_data_row:
        return
    for metric in METRIC_KEYS:
        if metric not in headers:
            continue
        col_idx = headers.index(metric) + 1
        values = [
            ws.cell(row_idx, col_idx).value
            for row_idx in range(first_data_row, last_data_row + 1)
            if isinstance(ws.cell(row_idx, col_idx).value, (int, float))
        ]
        if not values:
            continue
        best = min(values) if metric in LOWER_IS_BETTER else max(values)
        for row_idx in range(first_data_row, last_data_row + 1):
            cell = ws.cell(row_idx, col_idx)
            if cell.value == best:
                cell.font = bold
                cell.fill = fill


def set_widths(ws, widths: dict[int, int]) -> None:
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def format_numeric_columns(ws, *, first_row: int) -> None:
    for row in ws.iter_rows(min_row=first_row):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00000"


def row_sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return (
        row.get("crop", ""),
        index_or_end(MODEL_ORDER, row.get("model")),
        index_or_end(WINDOW_ORDER, row.get("window")),
        index_or_end(FEATURE_ORDER, row.get("feature_set")),
    )


def index_or_end(values: tuple[str, ...], value: object) -> int:
    try:
        return values.index(str(value))
    except ValueError:
        return len(values)


def average_key(rows: list[dict[str, Any]], key: str) -> float | None:
    values = [float(row[key]) for row in rows if row.get(key) is not None]
    if not values:
        return None
    return mean(values)


def label_value(row: dict[str, Any] | None, key: str) -> str | None:
    if row is None:
        return None
    return str(row.get(key))


def metric_value(row: dict[str, Any] | None, key: str) -> float | None:
    if row is None or row.get(key) is None:
        return None
    return float(row[key])


def to_optional_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


if __name__ == "__main__":
    main()

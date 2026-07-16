from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


WINDOW_LABELS = {
    "6mo": "6month",
    "9mo": "9month",
    "1year": "12month",
}
FEATURE_LABELS = {
    "ndvi_only": "NDVI",
    "all_indices": "NDVI, NDWI, EVI",
}
WINDOW_ORDER = ("6mo", "9mo", "1year")
FEATURE_ORDER = ("ndvi_only", "all_indices")
CROP_ORDER = ("Soybeans", "Rice", "Corn")
METRICS = (
    ("MAE_MEAN", "mae_mean", "min"),
    ("MAE_END", "mae_end", "min"),
    ("IOU_MEAN", "iou_mean", "max"),
    ("MAE_START", "mae_start", "min"),
    ("MAE_WINDOW", "mae_window", "min"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create an RNN harvest-window comparison workbook.")
    parser.add_argument("--results-root", type=Path, required=True)
    parser.add_argument("--out-xlsx", type=Path, required=True)
    parser.add_argument("--train-year", type=str, default="2022")
    parser.add_argument("--test-year", type=str, default="2023")
    parser.add_argument("--status-tsv", type=Path, default=None, help="Accepted for compatibility.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows_by_crop = collect_comparison_rows(args.results_root)
    args.out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    write_comparison_workbook(
        args.out_xlsx,
        rows_by_crop,
        train_year=args.train_year,
        test_year=args.test_year,
    )
    print(f"Wrote RNN comparison workbook to {args.out_xlsx}")


def collect_comparison_rows(results_root: Path) -> dict[str, list[dict[str, Any]]]:
    rows_by_crop: dict[str, list[dict[str, Any]]] = {}
    for crop in CROP_ORDER:
        crop_rows: list[dict[str, Any]] = []
        for window in WINDOW_ORDER:
            for feature_set in FEATURE_ORDER:
                metrics_path = results_root / window / feature_set / crop / "metrics.json"
                if not metrics_path.exists():
                    continue
                with open(metrics_path, "r", encoding="utf-8") as f:
                    metrics = json.load(f)
                test_metrics = metrics.get("test_metrics", {})
                row = {
                    "Model": f"{WINDOW_LABELS[window]} {FEATURE_LABELS[feature_set]}",
                }
                for header, key, _direction in METRICS:
                    row[header] = round(float(test_metrics[key]), 5)
                crop_rows.append(row)
        rows_by_crop[crop] = crop_rows
    return rows_by_crop


def write_comparison_workbook(
    path: Path,
    rows_by_crop: dict[str, list[dict[str, Any]]],
    *,
    train_year: str,
    test_year: str,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RNN"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    best_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.cell(1, 1, f"Train {train_year}").font = bold
    ws.cell(2, 1, f"Test {test_year}").font = bold

    row_idx = 4
    headers = ["Model"] + [header for header, _key, _direction in METRICS]
    for crop in CROP_ORDER:
        rows = rows_by_crop.get(crop, [])
        ws.cell(row_idx, 1, crop).font = bold
        row_idx += 2

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row_idx, col_idx, header)
            cell.font = bold
            cell.fill = header_fill
        header_row = row_idx
        row_idx += 1

        first_data_row = row_idx
        for row in rows:
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row_idx, col_idx, row.get(header))
                if col_idx > 1:
                    cell.number_format = "0.00000"
            row_idx += 1
        last_data_row = row_idx - 1

        bold_best_metric_values(
            ws,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
            first_metric_col=2,
            metrics=METRICS,
            bold=bold,
            fill=best_fill,
        )
        row_idx += 2 if crop == "Rice" else 1

    set_column_widths(ws)
    wb.save(path)


def bold_best_metric_values(
    ws,
    *,
    first_data_row: int,
    last_data_row: int,
    first_metric_col: int,
    metrics: tuple[tuple[str, str, str], ...],
    bold: Font,
    fill: PatternFill,
) -> None:
    if first_data_row > last_data_row:
        return
    for metric_offset, (_header, _key, direction) in enumerate(metrics):
        col_idx = first_metric_col + metric_offset
        values = [
            ws.cell(row_idx, col_idx).value
            for row_idx in range(first_data_row, last_data_row + 1)
            if isinstance(ws.cell(row_idx, col_idx).value, (int, float))
        ]
        if not values:
            continue
        best = max(values) if direction == "max" else min(values)
        for row_idx in range(first_data_row, last_data_row + 1):
            cell = ws.cell(row_idx, col_idx)
            if cell.value == best:
                cell.font = bold
                cell.fill = fill


def set_column_widths(ws) -> None:
    widths = {
        1: 24,
        2: 13,
        3: 13,
        4: 13,
        5: 13,
        6: 14,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


if __name__ == "__main__":
    main()

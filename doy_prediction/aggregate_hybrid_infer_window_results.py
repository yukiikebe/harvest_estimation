from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any, Tuple

import numpy as np
import torch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from doy_prediction.metrics import harvest_window_iou
from doy_prediction.predict_tile_hybrid_infer import (
    checkpoint_path,
    predict_cnn,
    predict_rnn,
    records_by_key,
)
from doy_prediction.tile_cnn_data import (
    TileCropRecord,
    build_tile_crop_records,
    load_crop_windows_yaml,
)


RecordKey = Tuple[int, str, str]

WINDOW_LABELS = {
    "6mo": "6month",
    "9mo": "9month",
    "1year": "12month",
}
FEATURE_LABELS = {
    "ndvi_only": "NDVI",
    "all_indices": "NDVI, NDWI, EVI",
}
WINDOW_ORDER = ("6mo", "9mo", "1year")
FEATURE_ORDER = ("ndvi_only", "all_indices")
CROP_ORDER = ("Soybeans", "Rice", "Corn")
METRICS = (
    ("Hybrid_infer_MAE_MEAN", "mae_mean", "min"),
    ("Hybrid_infer_MAE_END", "mae_end", "min"),
    ("Hybrid_infer_IOU_MEAN", "iou_mean", "max"),
    ("Hybrid_infer_MAE_START", "mae_start", "min"),
    ("Hybrid_infer_MAE_WINDOW", "mae_window", "min"),
)
CNN_CHECKPOINT_ROOTS = {
    "6mo": "harvest_cnn_first6mo_only2022",
    "9mo": "harvest_cnn_first9mo_only2022",
    "1year": "harvest_cnn_first_allmonth_only2022",
}
RNN_CHECKPOINT_ROOT = "harvest_rnn_window_sweep_32_64_128_cuda_2022_train_2023_test"
WINDOW_YAMLS = {
    "6mo": "configs/Arkansas/doy_input_first_6mo.yaml",
    "9mo": "configs/Arkansas/doy_input_first_9mo.yaml",
    "1year": None,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Evaluate late-fusion hybrid inference for each CNN/RNN window sweep row."
    )
    parser.add_argument("--inputs", type=Path, default=Path("outputs/2023_AR"))
    parser.add_argument("--models-root", type=Path, default=Path("outputs_prediction_DOY/models"))
    parser.add_argument("--out-csv", type=Path, default=Path("hybrid_infer_window_results.csv"))
    parser.add_argument("--out-xlsx", type=Path, default=Path("harvest Prediction Hybrid Infer.xlsx"))
    parser.add_argument("--device", type=str, default=None)
    parser.add_argument("--min-points", type=int, default=2)
    parser.add_argument("--train-year", type=str, default="2022")
    parser.add_argument("--test-year", type=str, default="2023")
    parser.add_argument(
        "--no-enforce-order",
        action="store_true",
        help="Do not force hybrid end DOY to be at least the hybrid start DOY.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    device = torch.device(args.device) if args.device else torch.device("cuda" if torch.cuda.is_available() else "cpu")
    rows = collect_rows(
        inputs=args.inputs,
        models_root=args.models_root,
        device=device,
        min_points=args.min_points,
        enforce_order=not args.no_enforce_order,
    )
    if not rows:
        raise RuntimeError("No hybrid inference rows were generated.")

    write_csv(args.out_csv, rows)
    write_workbook(
        args.out_xlsx,
        rows,
        train_year=args.train_year,
        test_year=args.test_year,
    )
    print(f"Wrote hybrid inference CSV to {args.out_csv}")
    print(f"Wrote hybrid inference workbook to {args.out_xlsx}")


def collect_rows(
    *,
    inputs: Path,
    models_root: Path,
    device: torch.device,
    min_points: int,
    enforce_order: bool,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for crop in CROP_ORDER:
        for window in WINDOW_ORDER:
            crop_window_yaml = WINDOW_YAMLS[window]
            crop_windows = load_crop_windows_yaml(Path(crop_window_yaml) if crop_window_yaml else None)
            for feature_set in FEATURE_ORDER:
                cnn_root = models_root / CNN_CHECKPOINT_ROOTS[window]
                rnn_root = models_root / RNN_CHECKPOINT_ROOT / window
                cnn_records = build_tile_crop_records(
                    inputs,
                    crops=[crop],
                    feature_set=feature_set,
                    min_points=min_points,
                    require_labels=True,
                    crop_windows=crop_windows,
                )
                rnn_records = build_tile_crop_records(
                    inputs,
                    crops=[crop],
                    feature_set=feature_set,
                    min_points=min_points,
                    require_labels=True,
                    crop_windows=crop_windows,
                )
                metrics = evaluate_combo(
                    crop=crop,
                    cnn_root=cnn_root,
                    rnn_root=rnn_root,
                    feature_set=feature_set,
                    cnn_records=cnn_records,
                    rnn_records=rnn_records,
                    device=device,
                    enforce_order=enforce_order,
                )
                row = {
                    "Crop": crop,
                    "Model": f"{WINDOW_LABELS[window]} {FEATURE_LABELS[feature_set]}",
                    "window": window,
                    "feature_set": feature_set,
                    "cnn_checkpoints": str(cnn_root),
                    "rnn_checkpoints": str(rnn_root),
                    **metrics,
                }
                rows.append(row)
    return rows


def evaluate_combo(
    *,
    crop: str,
    cnn_root: Path,
    rnn_root: Path,
    feature_set: str,
    cnn_records: list[TileCropRecord],
    rnn_records: list[TileCropRecord],
    device: torch.device,
    enforce_order: bool,
) -> dict[str, float]:
    if not cnn_records:
        raise RuntimeError(f"No CNN records found for crop={crop}, feature_set={feature_set}, root={cnn_root}")
    if not rnn_records:
        raise RuntimeError(f"No RNN records found for crop={crop}, feature_set={feature_set}, root={rnn_root}")

    cnn_by_key = records_by_key(cnn_records)
    rnn_by_key = records_by_key(rnn_records)
    common_keys = sorted(set(cnn_by_key) & set(rnn_by_key))
    if not common_keys:
        raise RuntimeError(f"No matching records for crop={crop}, feature_set={feature_set}")

    crop_cnn_records = [cnn_by_key[key] for key in common_keys]
    crop_rnn_records = [rnn_by_key[key] for key in common_keys]
    cnn_checkpoint = checkpoint_path(cnn_root, feature_set, crop)
    rnn_checkpoint = checkpoint_path(rnn_root, feature_set, crop)
    if not cnn_checkpoint.exists():
        raise FileNotFoundError(f"Missing CNN checkpoint: {cnn_checkpoint}")
    if not rnn_checkpoint.exists():
        raise FileNotFoundError(f"Missing RNN checkpoint: {rnn_checkpoint}")

    cnn_pred = predict_cnn(cnn_checkpoint, crop_cnn_records, device)
    rnn_pred = predict_rnn(rnn_checkpoint, crop_rnn_records, device)

    true_start = []
    true_end = []
    pred_start = []
    pred_end = []
    ious = []
    for key, cnn_record, rnn_record, cnn_row, rnn_row in zip(
        common_keys,
        crop_cnn_records,
        crop_rnn_records,
        cnn_pred,
        rnn_pred,
    ):
        if cnn_record.start_doy is None or cnn_record.end_doy is None:
            raise RuntimeError(f"Missing labels for CNN record: {key}")
        if rnn_record.start_doy is None or rnn_record.end_doy is None:
            raise RuntimeError(f"Missing labels for RNN record: {key}")
        start = int(cnn_row[0])
        end = int(rnn_row[1])
        if enforce_order:
            end = max(start, end)
        true_start.append(int(cnn_record.start_doy))
        true_end.append(int(cnn_record.end_doy))
        pred_start.append(start)
        pred_end.append(end)
        ious.append(harvest_window_iou(start, end, int(cnn_record.start_doy), int(cnn_record.end_doy)))

    true_start_arr = np.asarray(true_start, dtype=np.float32)
    true_end_arr = np.asarray(true_end, dtype=np.float32)
    pred_start_arr = np.asarray(pred_start, dtype=np.float32)
    pred_end_arr = np.asarray(pred_end, dtype=np.float32)
    mae_start = float(np.mean(np.abs(pred_start_arr - true_start_arr)))
    mae_end = float(np.mean(np.abs(pred_end_arr - true_end_arr)))
    return {
        "num_records": float(len(common_keys)),
        "mae_start": mae_start,
        "mae_end": mae_end,
        "mae_mean": (mae_start + mae_end) / 2.0,
        "mae_window": float(np.mean(np.abs((pred_end_arr - pred_start_arr) - (true_end_arr - true_start_arr)))),
        "iou_mean": float(np.mean(np.asarray(ious, dtype=np.float32))),
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "Crop",
        "Model",
        "Hybrid_infer_MAE_MEAN",
        "Hybrid_infer_MAE_END",
        "Hybrid_infer_IOU_MEAN",
        "Hybrid_infer_MAE_START",
        "Hybrid_infer_MAE_WINDOW",
        "num_records",
        "window",
        "feature_set",
        "cnn_checkpoints",
        "rnn_checkpoints",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(to_output_row(row))


def write_workbook(
    path: Path,
    rows: list[dict[str, Any]],
    *,
    train_year: str,
    test_year: str,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hybrid Infer"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    best_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.cell(1, 1, f"Train {train_year}").font = bold
    ws.cell(2, 1, f"Test {test_year}").font = bold
    headers = ["Model"] + [header for header, _key, _direction in METRICS]

    rows_by_crop = {crop: [row for row in rows if row["Crop"] == crop] for crop in CROP_ORDER}
    row_idx = 4
    for crop in CROP_ORDER:
        ws.cell(row_idx, 1, crop).font = bold
        row_idx += 2
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row_idx, col_idx, header)
            cell.font = bold
            cell.fill = header_fill
        row_idx += 1
        first_data_row = row_idx
        for row in rows_by_crop[crop]:
            output_row = to_output_row(row)
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row_idx, col_idx, output_row.get(header))
                if col_idx > 1:
                    cell.number_format = "0.00000"
            row_idx += 1
        last_data_row = row_idx - 1
        bold_best_metric_values(
            ws,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
            first_metric_col=2,
            metrics=METRICS,
            bold=bold,
            fill=best_fill,
        )
        row_idx += 2 if crop == "Rice" else 1

    set_column_widths(ws)
    wb.save(path)


def to_output_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        **row,
        "Hybrid_infer_MAE_MEAN": round(float(row["mae_mean"]), 5),
        "Hybrid_infer_MAE_END": round(float(row["mae_end"]), 5),
        "Hybrid_infer_IOU_MEAN": round(float(row["iou_mean"]), 5),
        "Hybrid_infer_MAE_START": round(float(row["mae_start"]), 5),
        "Hybrid_infer_MAE_WINDOW": round(float(row["mae_window"]), 5),
    }


def bold_best_metric_values(
    ws,
    *,
    first_data_row: int,
    last_data_row: int,
    first_metric_col: int,
    metrics: tuple[tuple[str, str, str], ...],
    bold: Font,
    fill: PatternFill,
) -> None:
    if first_data_row > last_data_row:
        return
    for metric_offset, (_header, _key, direction) in enumerate(metrics):
        col_idx = first_metric_col + metric_offset
        values = [
            ws.cell(row_idx, col_idx).value
            for row_idx in range(first_data_row, last_data_row + 1)
            if isinstance(ws.cell(row_idx, col_idx).value, (int, float))
        ]
        if not values:
            continue
        best = max(values) if direction == "max" else min(values)
        for row_idx in range(first_data_row, last_data_row + 1):
            cell = ws.cell(row_idx, col_idx)
            if cell.value == best:
                cell.font = bold
                cell.fill = fill


def set_column_widths(ws) -> None:
    widths = {
        1: 24,
        2: 22,
        3: 20,
        4: 21,
        5: 23,
        6: 24,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


if __name__ == "__main__":
    main()
